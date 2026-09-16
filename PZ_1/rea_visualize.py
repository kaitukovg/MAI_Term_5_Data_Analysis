import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from docx import Document
from docx.shared import Inches


df = pd.read_csv("rea_news.csv")

df["date"] = pd.to_datetime(df["date"])

df["month"] = df["date"].dt.month
df["year"] = df["date"].dt.year

sns.countplot(data=df, x="month", hue="month", palette="Blues_d")

plt.legend().remove()
plt.xlabel("Month")
plt.ylabel("News")
plt.title("News by months")
plt.xticks(rotation=30)
plt.savefig("rea_news.png", dpi=300)
plt.show()

plt.close()

doc = Document()

doc.add_heading(
    "Отчёт по публикациям РЭУ",
    level=1
)

doc.add_heading(
    "Распределение публикаций по месяцам",
    level=2
)

doc.add_picture(
    "rea_news.png",
    width=Inches(6)
)


doc.add_heading(
    "Публикации",
    level=2
)


columns = [
    col for col in df.columns
    if col not in ["month", "year"]
]

table = doc.add_table(
    rows=1,
    cols=len(columns)
)

table.style = "Table Grid"


for i, column in enumerate(columns):
    table.rows[0].cells[i].text = str(column)


for _, row in df.iterrows():

    cells = table.add_row().cells

    for i, column in enumerate(columns):

        value = row[column]

        if pd.isna(value):
            value = ""

        elif isinstance(value, pd.Timestamp):
            value = value.strftime("%Y-%m-%d")

        cells[i].text = str(value)


doc.save("rea_news_report.docx")

print("Отчёт создан: rea_news_report.docx")
